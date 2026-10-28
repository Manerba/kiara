"""Reproduzierbare Dokumentstichproben und eingefrorene Referenzen fuer Agent-Tests."""

from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from kiara.benchmarks.artifacts import write_private_json


def digest(value) -> str:
    return hashlib.sha256(
        json.dumps(value, ensure_ascii=False, sort_keys=True, default=str).encode()
    ).hexdigest()


def normalized(text: str) -> str:
    return " ".join(unicodedata.normalize("NFKC", text).casefold().split())


def extract_reference(path: Path) -> tuple[str, str]:
    """Liest Quelldateien ohne Suche, Ingestion oder Modellaufruf."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n\n".join(page.extract_text() or "" for page in PdfReader(path).pages), "pypdf"
    if suffix == ".docx":
        from docx import Document

        document = Document(path)
        paragraphs = [para.text for para in document.paragraphs]
        tables = [
            " | ".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        ]
        return "\n".join(paragraphs + tables), "python-docx"
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return "\n".join(
                sheet.title
                + "\n"
                + "\n".join(
                    " | ".join(str(value) if value is not None else "" for value in row)
                    for row in sheet.iter_rows(values_only=True)
                )
                for sheet in workbook
            ), "openpyxl"
        finally:
            workbook.close()
    raw = path.read_bytes()
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16"), "text:utf-16"
    try:
        return raw.decode("utf-8-sig"), "text:utf-8"
    except UnicodeDecodeError:
        from charset_normalizer import from_bytes

        decoded = from_bytes(raw).best()
        if decoded is None:
            raise ValueError("Referenz-Encoding nicht erkennbar") from None
        return str(decoded), f"text:{decoded.encoding}"


def sample_documents(
    *,
    quotas: dict[int, int],
    seed: str,
    embedding_model: str,
    output: Path,
    mount_map: dict[str, str] | None = None,
) -> dict:
    """Zieht Dateien nach DS-Quoten aus file_hashes, unabhaengig von Suchtreffern.

    mount_map ordnet private, bereits gemountete SMB-Roots den indexierten
    /mnt/kiara-ds-ID-Pfaden zu. Referenzpfade im Test bleiben unveraendert.
    """
    from kiara.shared.db import db_query

    if output.exists():
        raise FileExistsError(output)
    if not quotas or any(type(n) is not int or n < 1 for n in quotas.values()):
        raise ValueError("Positive Dokumentquoten erforderlich")
    excluded_pattern = (
        r"hetzner|apple|/node_modules/|/venv/|/\.venv/|/\.git/|/vendor/|"
        r"/dist/|/_internal/|package-lock|changelog|benchmark|canary|golden_set|"
        r"/tests?/|CLAUDE\.md|AGENTS\.md"
    )
    with db_query() as connection:
        candidates = connection.execute(
            """SELECT * FROM (
                SELECT f.datasource_id,d.name,f.file_path,f.file_size,
                    count(*) OVER (PARTITION BY f.datasource_id) AS population_size,
                    row_number() OVER (PARTITION BY f.datasource_id
                        ORDER BY md5(f.file_path || %s), f.file_path) AS sample_rank
                FROM file_hashes f JOIN datasources d ON d.id=f.datasource_id
                WHERE f.embedding_model=%s AND d.is_active
                    AND f.datasource_id=ANY(%s)
                    AND f.file_path ~* %s AND f.file_path !~* %s
                    AND f.file_size >= 500
            ) ranked WHERE sample_rank <= %s ORDER BY datasource_id,sample_rank""",
            (
                seed,
                embedding_model,
                list(quotas),
                r"\.(txt|md|pdf|docx|csv|xlsx|py|js|sh|html)$",
                excluded_pattern,
                max(quotas.values()) * 5,
            ),
        ).fetchall()
    counts = defaultdict(int)
    documents, excluded = [], []
    for candidate in candidates:
        row = dict(candidate)
        ds_id = row["datasource_id"]
        if counts[ds_id] >= quotas[ds_id]:
            continue
        source = Path(row["file_path"])
        local_path = source
        if str(ds_id) in (mount_map or {}):
            local_path = Path(mount_map[str(ds_id)]) / source.relative_to(f"/mnt/kiara-ds-{ds_id}")
        try:
            raw = local_path.read_bytes()
            text, extraction = extract_reference(local_path)
            if len(text.strip()) < 100:
                raise ValueError("Weniger als 100 extrahierte Zeichen")
        except Exception as exc:
            excluded.append({**row, "reason": f"{type(exc).__name__}: {exc}"})
            continue
        family = (
            f"project-ds{ds_id}"
            if row["name"].startswith("Projekt")
            else hashlib.sha256(source.stem.casefold().encode()).hexdigest()[:16]
        )
        documents.append(
            {
                **row,
                "id": f"D{len(documents) + 1:02d}",
                "family": family,
                "source": str(source),
                "reference_path": str(local_path),
                "source_sha256": hashlib.sha256(raw).hexdigest(),
                "text": text,
                "text_sha256": hashlib.sha256(text.encode()).hexdigest(),
                "extraction": extraction,
                "reference_layer": "mail_cache"
                if "/mail-cache/" in str(source)
                else "original_file",
            }
        )
        counts[ds_id] += 1
    families = sorted(
        {doc["family"] for doc in documents},
        key=lambda family: hashlib.sha256(f"{seed}:split:{family}".encode()).hexdigest(),
    )
    holdout = set(families[: max(1, len(families) // 3)])
    for document in documents:
        document["split"] = "holdout" if document["family"] in holdout else "development"
    result = {
        "seed": seed,
        "embedding_model": embedding_model,
        "quotas": quotas,
        "selection": "file_hashes, seeded MD5 order within each datasource; minimum 500 bytes",
        "exclusion_pattern": excluded_pattern,
        "excluded": excluded,
        "counts": dict(counts),
        "documents": documents,
    }
    write_private_json(output, result)
    return result


def validate_suite(suite: dict) -> None:
    documents = suite["documents"]
    by_id = {document["id"]: document for document in documents}
    if not documents or len(by_id) != len(documents):
        raise ValueError("Fehlende oder doppelte Dokument-IDs")
    families, contents = {}, {}
    for document in documents:
        split = document["split"]
        if split not in {"development", "holdout"}:
            raise ValueError("Ungueltige Teilmenge")
        for mapping, key in ((families, document["family"]), (contents, document["text_sha256"])):
            if mapping.setdefault(key, split) != split:
                raise ValueError("Dokumentfamilie oder identischer Inhalt in beiden Teilmengen")
        if hashlib.sha256(document["text"].encode()).hexdigest() != document["text_sha256"]:
            raise ValueError("Referenztext-Hash stimmt nicht")
    seen = set()
    for case in suite["cases"]:
        if case["id"] in seen or not re.fullmatch(r"[A-Za-z0-9_-]+", case["id"]):
            raise ValueError("Ungueltige oder doppelte Fall-ID")
        seen.add(case["id"])
        if type(case.get("repetitions", 1)) is not int or case.get("repetitions", 1) < 1:
            raise ValueError("Positive Wiederholungszahl erforderlich")
        if not str(case["question"]).strip() or case["split"] not in {"development", "holdout"}:
            raise ValueError("Frage und Teilmenge erforderlich")
        if case.get("expect_abstention"):
            if not case.get("negative_rationale") or case.get("evidence"):
                raise ValueError(
                    "Negativfall braucht Begruendung und darf keine Antwortbelege haben"
                )
        elif not case.get("evidence") or not case.get("reference_answer"):
            raise ValueError("Antwort und Belege erforderlich")
        for doc_id in case.get("scope_document_ids", []):
            if by_id[doc_id]["split"] != case["split"]:
                raise ValueError("Negativfall vermischt Entwicklungs- und Pruefdokumente")
        for evidence in case.get("evidence", []):
            document = by_id[evidence["document_id"]]
            if document["split"] != case["split"]:
                raise ValueError("Fall vermischt Entwicklungs- und Pruefdokumente")
            if not evidence.get("quotes") or any(
                not quote.strip() or normalized(quote) not in normalized(document["text"])
                for quote in evidence["quotes"]
            ):
                raise ValueError(f"{case['id']}: Beleg fehlt im Referenztext")
        for check in case.get("answer_checks", []):
            if not check or any(not str(value).strip() for value in check):
                raise ValueError("Leere Antwortpruefung")
    if not seen:
        raise ValueError("Keine Testfaelle")


def freeze_suite(draft: Path, output: Path) -> dict:
    if output.exists():
        raise FileExistsError(output)
    suite = json.loads(draft.read_text())
    suite.pop("sha256", None)
    validate_suite(suite)
    suite["frozen_at"] = datetime.now(timezone.utc).isoformat()
    suite["sha256"] = digest(suite)
    write_private_json(output, suite)
    return suite


def load_suite(path: Path) -> dict:
    suite = json.loads(path.read_text())
    expected = suite.pop("sha256")
    if digest(suite) != expected:
        raise ValueError("Die eingefrorene Testsuite wurde veraendert")
    validate_suite(suite)
    suite["sha256"] = expected
    return suite
