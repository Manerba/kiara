"""Dokumenten-Export fuer KIara.

Exportiert Markdown-Inhalte nach PDF, DOCX, reinem Text oder
belasst sie als Markdown-Datei.
"""

import logging
import os
import re
import time
from datetime import datetime, timezone
from pathlib import Path

logger = logging.getLogger(__name__)


class DocumentExporter:
    """Exportiert Dokumente in verschiedene Formate.

    Args:
        config: KIara-Konfiguration. Relevanter Schluessel:
                ``workflows.output_dir`` (default ``/var/lib/kiara/exports``).
    """

    def __init__(self, config: dict) -> None:
        from kiara.config import get_export_dir

        self._output_dir = get_export_dir(config)

    @property
    def output_dir(self) -> str:
        """Public accessor for the export output directory."""
        return self._output_dir

    # ------------------------------------------------------------------
    # Oeffentliche API
    # ------------------------------------------------------------------

    async def export(
        self,
        content: str,
        filename: str,
        format: str = "markdown",
        title: str | None = None,
        metadata: dict | None = None,
    ) -> str:
        """Exportiert Inhalt in das gewuenschte Format.

        Args:
            content: Markdown-Inhalt.
            filename: Dateiname (ohne Erweiterung wird angepasst).
            format: Zielformat (``markdown``, ``pdf``, ``docx``, ``text``).
            title: Optionaler Dokumenttitel.
            metadata: Optionale Metadaten.

        Returns:
            Absoluter Pfad zur exportierten Datei.
        """
        out_dir = Path(self._output_dir)
        out_dir.mkdir(parents=True, exist_ok=True)

        # Dateiname bereinigen
        safe_name = self._sanitize_filename(filename)

        # Erweiterung anpassen
        ext_map = {
            "markdown": ".md",
            "pdf": ".pdf",
            "docx": ".docx",
            "text": ".txt",
        }
        if format not in ext_map:
            raise ValueError(
                f"Unbekanntes Exportformat: {format!r}. "
                f"Erlaubte Formate: {', '.join(sorted(ext_map))}"
            )
        expected_ext = ext_map[format]
        # Mehrfach-Extensions wie .tar.gz erhalten: nur die genau gleiche Zielendung
        # wird abgeschnitten, sonst wird expected_ext einfach angehaengt.
        if safe_name.lower().endswith(expected_ext):
            base_name = safe_name[: -len(expected_ext)]
        else:
            base_name = safe_name
        filepath = out_dir / f"{base_name}{expected_ext}"

        # Eindeutigen Dateinamen erzeugen wenn Datei existiert oder ein Symlink vorhanden ist.
        # is_symlink() ist symlink-honest (auch broken Symlinks) — verhindert Schreibvorgaenge
        # ueber praeparierte Symlinks ausserhalb von out_dir.
        counter = 1
        while filepath.exists() or filepath.is_symlink():
            filepath = out_dir / f"{base_name}_{counter}{expected_ext}"
            counter += 1

        try:
            if format == "pdf":
                await self._export_pdf(content, filepath, title, metadata)
            elif format == "docx":
                await self._export_docx(content, filepath, title, metadata)
            elif format == "text":
                await self._export_text(content, filepath)
            else:
                # Markdown: direkt schreiben
                filepath.write_text(content, encoding="utf-8")
        except Exception:
            # Halb-geschriebene Datei aufraeumen, damit list_exports keine
            # orphan-Eintraege liefert. Exception fliegt unveraendert weiter.
            filepath.unlink(missing_ok=True)
            raise

        logger.info("Dokument exportiert: %s (%s)", filepath, format)
        return str(filepath.resolve())

    def list_exports(self, limit: int = 50) -> list[dict]:
        """Listet exportierte Dateien auf.

        Args:
            limit: Maximale Anzahl zurueckgegebener Eintraege.

        Returns:
            Liste von Dicts mit ``filename``, ``format``, ``size_bytes``,
            ``created_at``.
        """
        out_dir = Path(self._output_dir)
        if not out_dir.exists():
            return []

        format_map = {
            ".md": "markdown",
            ".pdf": "pdf",
            ".docx": "docx",
            ".txt": "text",
        }

        # Snapshot mit stat-Daten sammeln, bevor sortiert wird. Race mit
        # cleanup_exports: Datei kann zwischen iterdir() und stat() verschwinden.
        entries_with_stats: list[tuple[Path, os.stat_result]] = []
        for entry in out_dir.iterdir():
            try:
                if not entry.is_file():
                    continue
                entries_with_stats.append((entry, entry.stat()))
            except FileNotFoundError:
                continue

        entries_with_stats.sort(key=lambda pair: pair[1].st_mtime, reverse=True)

        exports = []
        for entry, stat in entries_with_stats:
            ext = entry.suffix.lower()
            exports.append(
                {
                    "filename": entry.name,
                    "format": format_map.get(ext, ext),
                    "size_bytes": stat.st_size,
                    "created_at": datetime.fromtimestamp(
                        stat.st_mtime, tz=timezone.utc
                    ).isoformat(),
                }
            )
            if len(exports) >= limit:
                break

        return exports

    def cleanup_exports(self, max_age_days: int = 30) -> int:
        """Loescht Exports die aelter als ``max_age_days`` Tage sind.

        Args:
            max_age_days: Maximales Alter in Tagen.

        Returns:
            Anzahl geloeschter Dateien.
        """
        out_dir = Path(self._output_dir)
        if not out_dir.exists():
            return 0

        cutoff = time.time() - (max_age_days * 86400)
        deleted = 0

        for entry in out_dir.iterdir():
            if not entry.is_file():
                continue
            if entry.stat().st_mtime < cutoff:
                try:
                    entry.unlink()
                    deleted += 1
                    logger.info("Export geloescht (zu alt): %s", entry.name)
                except OSError as exc:
                    logger.warning("Export konnte nicht geloescht werden: %s (%s)", entry.name, exc)

        return deleted

    # ------------------------------------------------------------------
    # Format-spezifische Exporter
    # ------------------------------------------------------------------

    async def _export_pdf(
        self,
        content: str,
        filepath: Path,
        title: str | None = None,
        metadata: dict | None = None,
    ) -> None:
        """Exportiert Markdown-Inhalt als PDF via reportlab.

        Unterstuetzt Ueberschriften (#, ##, ###), Listen (- ),
        Fettschrift (**text**) und normalen Text.

        Args:
            content: Markdown-Inhalt.
            filepath: Zielpfad fuer die PDF-Datei.
            title: Optionaler Dokumenttitel.
            metadata: Optionale Metadaten.
        """
        try:
            from reportlab.lib.colors import HexColor
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
            from reportlab.lib.units import mm
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        except ImportError as exc:
            raise ValueError("PDF-Export benoetigt das Paket 'reportlab'.") from exc

        doc = SimpleDocTemplate(
            str(filepath),
            pagesize=A4,
            topMargin=25 * mm,
            bottomMargin=25 * mm,
            leftMargin=20 * mm,
            rightMargin=20 * mm,
        )

        styles = getSampleStyleSheet()
        accent_color = HexColor("#2d5f8a")

        # Eigene Styles
        style_h1 = ParagraphStyle(
            "CustomH1",
            parent=styles["Heading1"],
            textColor=accent_color,
            fontSize=18,
            spaceAfter=12,
            spaceBefore=16,
        )
        style_h2 = ParagraphStyle(
            "CustomH2",
            parent=styles["Heading2"],
            textColor=accent_color,
            fontSize=14,
            spaceAfter=8,
            spaceBefore=12,
        )
        style_h3 = ParagraphStyle(
            "CustomH3",
            parent=styles["Heading3"],
            textColor=accent_color,
            fontSize=12,
            spaceAfter=6,
            spaceBefore=10,
        )
        style_body = styles["Normal"]
        style_list = ParagraphStyle(
            "CustomList",
            parent=styles["Normal"],
            leftIndent=15,
            bulletIndent=5,
            spaceBefore=2,
            spaceAfter=2,
        )

        story: list = []

        # Titel
        if title:
            story.append(Paragraph(self._escape_xml(title), style_h1))
            story.append(Spacer(1, 6))

        # Metadaten
        if metadata:
            meta_parts = []
            for key, value in metadata.items():
                meta_parts.append(
                    f"<b>{self._escape_xml(str(key))}:</b> {self._escape_xml(str(value))}"
                )
            meta_text = " | ".join(meta_parts)
            story.append(Paragraph(meta_text, styles["Italic"]))
            story.append(Spacer(1, 10))

        # Inhalt parsen
        for line in content.split("\n"):
            stripped = line.strip()

            if not stripped:
                story.append(Spacer(1, 6))
                continue

            if stripped.startswith("### "):
                text = self._md_to_reportlab(stripped[4:])
                story.append(Paragraph(text, style_h3))
            elif stripped.startswith("## "):
                text = self._md_to_reportlab(stripped[3:])
                story.append(Paragraph(text, style_h2))
            elif stripped.startswith("# "):
                text = self._md_to_reportlab(stripped[2:])
                story.append(Paragraph(text, style_h1))
            elif stripped[:6].lower() == "- [ ] " or stripped[:6].lower() == "- [x] ":
                checkbox = "\u2610 " if stripped[:5].lower() == "- [ ]" else "\u2611 "
                text = self._md_to_reportlab(stripped[6:])
                story.append(Paragraph(f"{checkbox}{text}", style_list))
            elif stripped.startswith("- "):
                text = self._md_to_reportlab(stripped[2:])
                story.append(Paragraph(f"\u2022 {text}", style_list))
            elif stripped.startswith("---"):
                story.append(Spacer(1, 10))
            else:
                text = self._md_to_reportlab(stripped)
                story.append(Paragraph(text, style_body))

        doc.build(story)

    async def _export_docx(
        self,
        content: str,
        filepath: Path,
        title: str | None = None,
        metadata: dict | None = None,
    ) -> None:
        """Exportiert Markdown-Inhalt als DOCX via python-docx.

        Mapping:
        - ``#``  -> Heading 1
        - ``##`` -> Heading 2
        - ``###``-> Heading 3
        - ``- `` -> List Bullet
        - Rest   -> Normal

        Args:
            content: Markdown-Inhalt.
            filepath: Zielpfad fuer die DOCX-Datei.
            title: Optionaler Dokumenttitel.
            metadata: Optionale Metadaten.
        """
        try:
            from docx import Document
        except ImportError:
            raise ValueError("python-docx ist nicht installiert")

        doc = Document()

        if title:
            doc.add_heading(title, level=0)

        if metadata:
            meta_parts = []
            for key, value in metadata.items():
                meta_parts.append(f"{key}: {value}")
            meta_para = doc.add_paragraph()
            meta_run = meta_para.add_run(" | ".join(meta_parts))
            meta_run.italic = True

        for line in content.split("\n"):
            stripped = line.strip()

            if not stripped:
                continue

            if stripped.startswith("### "):
                text = self._strip_md_formatting(stripped[4:])
                doc.add_heading(text, level=3)
            elif stripped.startswith("## "):
                text = self._strip_md_formatting(stripped[3:])
                doc.add_heading(text, level=2)
            elif stripped.startswith("# "):
                text = self._strip_md_formatting(stripped[2:])
                doc.add_heading(text, level=1)
            elif stripped[:6].lower() == "- [ ] " or stripped[:6].lower() == "- [x] ":
                text = self._strip_md_formatting(stripped[6:])
                prefix = "\u2610 " if stripped[:5].lower() == "- [ ]" else "\u2611 "
                doc.add_paragraph(f"{prefix}{text}", style="List Bullet")
            elif stripped.startswith("- "):
                text = self._strip_md_formatting(stripped[2:])
                doc.add_paragraph(text, style="List Bullet")
            elif stripped.startswith("---"):
                continue
            else:
                text = self._strip_md_formatting(stripped)
                doc.add_paragraph(text)

        doc.save(str(filepath))

    async def _export_text(self, content: str, filepath: Path) -> None:
        """Exportiert als reinen Text (Markdown-Markup entfernt).

        Args:
            content: Markdown-Inhalt.
            filepath: Zielpfad fuer die Textdatei.
        """
        clean = self._strip_md_formatting(content)
        filepath.write_text(clean, encoding="utf-8")

    # ------------------------------------------------------------------
    # Hilfsmethoden
    # ------------------------------------------------------------------

    @staticmethod
    def _sanitize_filename(name: str) -> str:
        """Bereinigt einen Dateinamen (entfernt unsichere Zeichen)."""
        # Nur Buchstaben, Zahlen, Bindestrich, Unterstrich, Punkt
        safe = re.sub(r"[^\w\-.]", "_", name)
        # Doppelte Unterstriche und fuehrende Punkte entfernen
        safe = re.sub(r"_+", "_", safe).lstrip(".")
        # Auf UTF-8-Byte-Limit kuerzen: Dateisystem-Limit ist 255 Bytes pro
        # Pfadkomponente, nicht 255 Codepoints. Reserve fuer Extension (.docx)
        # plus Counter (_99999) belassen. errors="ignore" kappt halbierte
        # Multibyte-Sequenzen am Ende sauber.
        max_bytes = 200
        encoded = safe.encode("utf-8")
        if len(encoded) > max_bytes:
            safe = encoded[:max_bytes].decode("utf-8", errors="ignore")
        return safe or "export"

    @staticmethod
    def _strip_md_formatting(text: str) -> str:
        """Entfernt Markdown-Formatierung aus Text.

        Entfernt: ``**bold**``, ``# headings``, ``- lists``,
        ``[ ] checkboxes``, ``---`` Trennlinien.
        """
        # Fett (**text**)
        result = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
        # Kursiv (*text*), aber keine Rechen-/Wildcard-Sterne wie "5 * 3 * 2"
        result = re.sub(r"(?<!\*)\*([^\s*\n](?:[^*\n]*?[^\s*\n])?)\*(?!\*)", r"\1", result)
        # Ueberschriften am Zeilenanfang
        result = re.sub(r"^#{1,6}\s+", "", result, flags=re.MULTILINE)
        # Checkboxen
        result = re.sub(r"- \[[ x]\] ", "", result)
        # Listenzeichen
        result = re.sub(r"^- ", "", result, flags=re.MULTILINE)
        # Trennlinien
        result = re.sub(r"^-{3,}$", "", result, flags=re.MULTILINE)
        return result

    @staticmethod
    def _md_to_reportlab(text: str) -> str:
        """Konvertiert einfache Markdown-Formatierung in ReportLab-XML.

        ``**text**`` -> ``<b>text</b>``
        Escaped XML-Sonderzeichen im restlichen Text.
        """
        # Erst XML-Escaping, dann Markdown → XML-Tags
        escaped = (
            text.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace('"', "&quot;")
        )
        result = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", escaped)
        result = re.sub(r"\*(.+?)\*", r"<i>\1</i>", result)
        return result

    @staticmethod
    def _escape_xml(text: str) -> str:
        """Escaped XML-Sonderzeichen fuer ReportLab."""
        return (
            text.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace('"', "&quot;")
        )
