"""Allgemeine Dokument-Loader (Docx, PDF, XLSX, ODT, PPTX, CSV, JSON, Plaintext, Code)."""

import csv
import io
import json
import logging
from pathlib import Path

logger = logging.getLogger(__name__)


_CODE_LANGUAGES = {
    ".cs": "csharp",
    ".vb": "vb",
    ".ps1": "powershell",
    ".py": "python",
    ".pl": "perl",
    ".js": "javascript",
    ".bat": "batch",
    ".sh": "bash",
    ".css": "css",
    ".html": "html",
    ".sln": "text",
}


def _load_plaintext(path: str) -> str:
    """Load plain text files with automatic encoding detection."""
    try:
        return Path(path).read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        pass
    raw = Path(path).read_bytes()
    # CP1252 heuristic: bytes 0x80-0x9F are printable in CP1252 (€ „ " " – —)
    # but C1 controls in ISO-8859-1. charset_normalizer can misidentify short
    # German texts containing these bytes.
    if any(0x80 <= b <= 0x9F for b in raw):
        try:
            return raw.decode("cp1252")
        except (UnicodeDecodeError, LookupError):
            pass
    try:
        from charset_normalizer import from_bytes

        result = from_bytes(raw)
        best = result.best()
        if best is not None:
            return str(best)
    except (ImportError, OSError, ValueError, LookupError, TypeError):
        pass
    text = raw.decode("utf-8", errors="replace")
    if "\ufffd" in text:
        count = text.count("\ufffd")
        logger.warning(
            "Encoding-Fallback fuer %s: %d Zeichen durch U+FFFD ersetzt"
            " (charset_normalizer fehlgeschlagen, UTF-8 errors='replace' aktiv)",
            path,
            count,
        )
    return text


def _load_code(path: str) -> str:
    """Load code files with language annotation."""
    ext = Path(path).suffix.lower()
    lang = _CODE_LANGUAGES.get(ext, "")
    content = _load_plaintext(path)
    if lang:
        return f"```{lang}\n# File: {Path(path).name}\n{content}\n```"
    return content


def _load_docx(path: str) -> str:
    """Load .docx files."""
    from docx import Document

    with open(path, "rb") as f:
        doc = Document(f)
        paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))
        return "\n\n".join(paragraphs)


def _load_pdf(path: str) -> str:
    """Load .pdf files."""
    from pypdf import PdfReader

    with open(path, "rb") as f:
        reader = PdfReader(f)
        pages = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                pages.append(text.replace("\x00", ""))  # pypdf Encoding-Artefakte
        return "\n\n".join(pages)


def _load_xlsx(path: str) -> str:
    """Load .xlsx/.xlsm files."""
    from openpyxl import load_workbook

    wb = None
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        sheets = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                if any(c.strip() for c in cells):
                    rows.append(" | ".join(cells))
            if rows:
                sheets.append(f"## Sheet: {sheet_name}\n" + "\n".join(rows))
        return "\n\n".join(sheets)
    finally:
        if wb is not None:
            wb.close()


def _load_odt(path: str) -> str:
    """Load .odt files."""
    from odf import text as odf_text
    from odf.opendocument import load

    doc = load(path)
    paragraphs = []
    for p in doc.getElementsByType(odf_text.P):
        content = ""
        for node in p.childNodes:
            if hasattr(node, "data"):
                content += node.data
            else:
                content += str(node)
        if content.strip():
            paragraphs.append(content)
    return "\n\n".join(paragraphs)


def _load_pptx(path: str) -> str:
    """Load .pptx files."""
    from pptx import Presentation

    with open(path, "rb") as f:
        prs = Presentation(f)
        slides = []
        for i, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            texts.append(text)
            if texts:
                slides.append(f"## Slide {i}\n" + "\n".join(texts))
    return "\n\n".join(slides)


def _load_csv_file(path: str) -> str:
    """Load .csv files as readable text."""
    content = _load_plaintext(path)
    try:
        reader = csv.reader(io.StringIO(content))
        rows = []
        for row in reader:
            rows.append(" | ".join(row))
        return "\n".join(rows)
    except csv.Error:
        return content


def _load_json_file(path: str) -> str:
    """Load .json files with pretty formatting."""
    content = _load_plaintext(path)
    try:
        data = json.loads(content)
        return json.dumps(data, indent=2, ensure_ascii=False)
    except json.JSONDecodeError:
        return content
